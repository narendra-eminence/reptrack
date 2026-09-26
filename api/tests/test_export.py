from typing import cast

from conftest import SEARCH_BODY, FakeSearch, make_client, wait_until
from fastapi import FastAPI
from openpyxl import load_workbook

from pipeline_api.export import slug, write_rows_xlsx


def test_formula_like_text_stays_text_and_control_chars_are_stripped(tmp_path):
    path = tmp_path / "x.xlsx"
    write_rows_xlsx(path, ["Title", "Snippet"], [['=HYPERLINK("http://evil")', "bad\x0bchars\x1fhere"]], "Sheet")
    ws = load_workbook(path)["Sheet"]
    assert ws["A2"].value == '=HYPERLINK("http://evil")' and ws["A2"].data_type == "s"
    assert ws["B2"].value == "badcharshere"


def test_serp_export_columns_and_filename(settings, monkeypatch):
    monkeypatch.setenv("SERPAPI_KEY", "k")
    with make_client(settings, search_one=FakeSearch()) as c:
        run_id = c.post("/api/runs", json={**SEARCH_BODY, "confirmed_calls": 2}).json()["id"]
        wait_until(lambda: c.get(f"/api/runs/{run_id}").json()["status"] == "scraped")
        r = c.get(f"/api/runs/{run_id}/serp.xlsx")
        bs = cast(FastAPI, c.app).state.deps.bs
    assert r.status_code == 200
    assert 'filename="alpha_2026-03-01_2026-08-31_serp.xlsx"' in r.headers["content-disposition"]
    path = settings.exports_dir / run_id / "alpha_2026-03-01_2026-08-31_serp.xlsx"
    ws = load_workbook(path).active
    assert ws is not None
    assert [cell.value for cell in ws[1]] == bs.EXPORT_COLUMNS
    assert ws.max_row == 3


def test_slug():
    assert slug('"Safari" luggage OR bags / 2026') == "Safari-luggage-OR-bags-2026"
    assert slug("   ") == "run"
