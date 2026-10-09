"""The cleaning engine end to end on synthetic verified rows."""

import datetime as dt

import pytest
from openpyxl import load_workbook

from pipeline_api.clean.details import CleaningDetails
from pipeline_api.clean.engine import CleanContext, clean
from pipeline_api.clean.media import load_master_list
from pipeline_api.clean.workbook import write_workbook
from pipeline_api.settings import load_settings

RAW = ["Query", "Published", "Outlet", "Title", "Snippet", "Link"]


def row(link, title=None, snippet=None, query="acme", published="2026-09-22", outlet="", **extra):
    """A verified row. Title and snippet default to text unique to the link, so only the rows a test means to be
    the same story are."""
    return {
        "Query": query,
        "Published": published,
        "Outlet": outlet,
        "Title": title if title is not None else f"Acme news from {link}",
        "Snippet": snippet if snippet is not None else f"Acme was mentioned on {link} today",
        "Link": link,
        "Provider": "serpapi",
        "Status": "Verified",
        **extra,
    }


@pytest.fixture(scope="module")
def media():
    return load_master_list(load_settings({}).master_media_list)


def run_clean(rows, media, details=None):
    ctx = CleanContext(
        run_name="Acme run",
        brand_set="acme",
        verify_job_id=1,
        start=dt.date(2026, 9, 1),
        end=dt.date(2026, 9, 30),
        region="in",
        media=media,
        media_list_name="list.xlsx",
        details=details
        or CleaningDetails.from_dict(
            {"own_websites": ["acme.com"], "own_handles": ["acmebags"], "competitor_handles": ["rivalbags"]}
        ),
        mentions_brand=lambda text: "acme" in text.lower(),
        raw_columns=RAW,
    )
    return clean(rows, ctx)


def links(result, sheet):
    s = result.sheet(sheet)
    assert s is not None, sheet
    return [r["Link"] for r in s.rows]


ET = "https://economictimes.indiatimes.com/industry/acme-results/articleshow/1.cms"
MINT = "https://www.livemint.com/companies/acme-results-1.html"
WIRE = "Acme posts its best quarter on record as bag sales climb"


@pytest.fixture(scope="module")
def result(media):
    rows = [
        row(ET, title=WIRE, outlet="Economic Times"),
        row(ET + "?utm_source=twitter", title=WIRE, query="acme results"),
        row(MINT, title=WIRE, snippet="Mint: " + WIRE),
        row("https://unknown-news-site.example/acme-quarter", title=WIRE, snippet="copy"),
        row("https://www.acme.com/news/launch"),
        row("https://acme.com/news/launch/"),
        row("https://x.com/acmebags/status/11"),
        row("https://x.com/rivalbags/status/12", title="Rival takes on Acme with a cheaper trolley range"),
        row("https://www.amazon.in/dp/B0X"),
        row("https://www.amazon.in/dp/B0X?tag=x"),
        row("https://www.amazon.in/dp/B0X/"),
        row("https://pacess.in/expert-time/Acme-38-1"),
        row(
            "https://another-news.example/markets-today",
            title="Markets end higher on Friday",
            snippet="Sensex rose 300 points led by banks",
        ),
        row("https://old-news.example/acme-old", published="2026-08-01"),
        row(
            "https://www.instagram.com/p/DdjT1BxTINs/",
            published="",
            Author="travelwithmeVerified",
            title="Travel day with my Acme trolley",
            snippet="Packed my Acme trolley for Goa",
        ),
        row("https://page-date.example/acme-story", published="", **{"Published Date": "2026-09-25"}),
        row(
            "https://lead.example/acme",
            title="Acme opens a store in Pune",
            snippet="Acme opened a flagship store in Pune",
            **{"Opening Text": "Acme opened a flagship store in Pune"},
        ),
        row(
            "https://uk.linkedin.com/in/someone",
            title="Someone - Sales Manager at Acme",
            snippet="Sales Manager at Acme. Experience: 5 years",
        ),
        row("https://empty.example/x", title="", snippet="🎉🎉"),
    ]
    return run_clean(rows, media)


def test_sheet_order(result):
    names = [s.name for s in result.sheets]
    assert names[:2] == ["Cleaning Summary", "Clean Data"]
    assert names[2:12] == [
        "Major Media",
        "Regional Media",
        "Other Media",
        "Twitter",
        "YouTube",
        "Facebook",
        "Instagram",
        "Reddit",
        "LinkedIn",
        "Other Sources",
    ]
    assert names[12:18] == [
        "Brand Communication",
        "Low Quality",
        "Spam",
        "Unclassified",
        "Verification Removed",
        "Out of Range",
    ]
    assert names[-4:] == ["Source Bucket", "NEW Domains", "Query Yield", "Raw Data"]


def test_major_media_and_dedupe(result):
    major = result.sheet("Major Media").rows
    assert [r["Source Name"] for r in major] == ["Economic Times", "Mint Companies"]  # both Core: exempt
    assert all(r["Media Level"] == "Core" and r["Verification Status"] == "Pending" for r in major)
    assert major[0]["Query Count"] == 2
    assert links(result, "Major Media_duplicate") == [ET + "?utm_source=twitter"]
    assert result.sheet("Major Media_duplicate").rows[0]["Notes"] == "Duplicate link"
    other_dups = result.sheet("Other Media_duplicate").rows
    assert [r["Link"] for r in other_dups] == ["https://unknown-news-site.example/acme-quarter"]
    assert other_dups[0]["Notes"].startswith("Duplicate title/hit sentence of ")
    assert result.facts["core_saved"] == 1


def test_brand_communication_and_competitors(result):
    bc = result.sheet("Brand Communication").rows
    assert [r["Link"] for r in bc] == [
        "https://www.acme.com/news/launch",
        "https://acme.com/news/launch/",
        "https://x.com/acmebags/status/11",
    ]
    assert bc[1]["Notes"] == "Duplicate of https://www.acme.com/news/launch"
    assert {r["Source Bucket"] for r in bc} == {"Brand Communication"}
    assert {r["Verification Status"] for r in bc} == {"Not verified"}
    other = result.sheet("Other Media").rows
    assert other[-1]["Source Bucket"] == "Competitor Owned"
    clean_links = links(result, "Clean Data")
    assert "https://x.com/rivalbags/status/12" not in clean_links
    assert not any("acme.com" in link for link in clean_links)


def test_exclusions(result):
    low = {r["Link"]: r for r in result.sheet("Low Quality").rows}
    assert low["https://www.amazon.in/dp/B0X"]["Exclusion Type"] == "E-commerce"
    assert low["https://www.amazon.in/dp/B0X"]["Source Bucket"] == ""
    assert low["https://www.amazon.in/dp/B0X"]["Status"] == ""  # not a NEW media domain for the master list
    assert "https://www.amazon.in/dp/B0X/" not in low  # a copy of an excluded row is counted, not listed
    nbm = low["https://another-news.example/markets-today"]
    assert (nbm["Exclusion Type"], nbm["Source Bucket"]) == ("No brand mention", "Other Media")
    assert low["https://uk.linkedin.com/in/someone"]["Exclusion Type"] == "Profile / bio match"
    assert links(result, "Spam") == ["https://pacess.in/expert-time/Acme-38-1"]
    assert links(result, "Out of Range") == ["https://old-news.example/acme-old"]
    assert links(result, "Unclassified") == ["https://empty.example/x"]
    assert result.sheet("Verification Removed").rows == []


def test_dates_authors_and_opening_text(result):
    by_link = {r["Link"]: r for r in result.sheet("Clean Data").rows}
    insta = by_link["https://www.instagram.com/p/DdjT1BxTINs/"]
    assert (insta["Date"], insta["Date Source"], insta["Author"]) == (None, "missing", "travelwithme")
    page = by_link["https://page-date.example/acme-story"]
    assert (page["Date"], page["Date Source"]) == (dt.date(2026, 9, 25), "page")
    assert by_link["https://lead.example/acme"]["Opening Text"] == ""
    assert by_link[ET]["Date Source"] == "export"
    assert by_link[ET]["Author"] == "Economic Times"
    assert all(not by_link[ET].get(c) for c in ("Driver", "Sub Parameter", "Sentiment"))


def test_process_sheets_and_facts(result):
    raw = result.sheet("Raw Data")
    assert raw.columns == [*RAW, "Cleaning Outcome"]
    assert len(raw.rows) == 19
    outcomes = {r["Link"]: r["Cleaning Outcome"] for r in raw.rows}
    assert outcomes["https://www.amazon.in/dp/B0X/"] == "Not listed: duplicate link of a row on Low Quality"
    assert outcomes["https://x.com/rivalbags/status/12"] == "Other Media (Competitor Owned)"
    new = {r["Domain"] for r in result.sheet("NEW Domains").rows}
    assert "page-date.example" in new and "economictimes.indiatimes.com" not in new
    qy = {r["Query"]: r for r in result.sheet("Query Yield").rows}
    assert qy["acme results"]["Duplicates"] == 1
    f = result.facts
    assert f["rows_in"] == 19
    assert f["checks_ok"] is True
    assert f["brand_communication"] == 2
    assert f["competitor_owned"] == 1
    assert f["buckets"]["Major Media"] == 2


def test_workbook_round_trip(result, tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(path, result)
    wb = load_workbook(path)
    assert wb.sheetnames == [s.name for s in result.sheets]
    ws = wb["Clean Data"]
    header = [c.value for c in ws[1]]
    assert header[:15] == [
        "Date",
        "Link",
        "URL",
        "Author",
        "Title",
        "Opening Text",
        "Hit Sentence",
        "Driver",
        "Sub Parameter",
        "Sentiment",
        "Language",
        "Country",
        "Engagement",
        "Reach",
        "Source Bucket",
    ]
    assert ws.freeze_panes == "A2"
    assert wb["Low Quality"]["P1"].value == "Exclusion Type"
    summary = {r[0].value: r[1].value for r in wb["Cleaning Summary"].iter_rows(min_row=2)}
    assert summary["Rows in"] == 19
    assert summary["Source Bucket values allowed"] == "OK"


def test_empty_input(media):
    result = run_clean([], media, CleaningDetails())
    assert result.facts["rows_in"] == 0
    assert links(result, "Clean Data") == []
