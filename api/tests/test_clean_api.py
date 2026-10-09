"""Cleaning over the API: brand-set details, the clean job on a finished verification, retry and the workbook."""

import io

import pandas as pd
import pytest
from conftest import SEARCH_BODY, FakeSearch, make_client, wait_until
from openpyxl import load_workbook

LINKS = [
    "https://economictimes.indiatimes.com/industry/acme-q2/articleshow/1.cms",
    "https://www.acme-bags.com/press/q2",
    "https://www.amazon.in/dp/ACME1",
    "https://x.com/acmebags/status/1",
    "https://www.instagram.com/p/AcmePost1/",
]


class AcmeSearch:
    """FakeSearch rows whose title and snippet name Acme, so the brand-mention rule keeps them."""

    def __init__(self, results):
        self.inner = FakeSearch(results)

    def __call__(self, query, start, end, pages, vertical, provider, stop, region=None):
        rows, error, attempts = self.inner(query, start, end, pages, vertical, provider, stop, region)
        for i, r in enumerate(rows):
            r.update(
                title=f"Acme story number {i} for the quarter",
                snippet=f"Acme said story {i} is out now",
                published="2026-05-02",
                domain="",
                outlet="",
            )
        return rows, error, attempts


async def fake_pipeline(input_path, output_path, cfg, brand, cache_dir="cache", rules=None, on_result=None, **kw):
    df = pd.read_excel(input_path)
    for _ in range(len(df)):
        assert on_result is not None
        on_result("Verified")
    rows = [{**r, "Status": "Verified", "Duplicate Of Row": None, "HTTP Status": 200} for r in df.to_dict("records")]
    pd.DataFrame(rows).to_excel(output_path, index=False)
    return {"Verified": len(df), "_rows": len(df), "_duplicates": 0, "_unique": len(df)}


@pytest.fixture
def keys(monkeypatch):
    monkeypatch.setenv("SERPAPI_KEY", "k")


def _verified_run(c):
    run_id = c.post("/api/runs", json={**SEARCH_BODY, "queries": "alpha", "confirmed_calls": 1}).json()["id"]
    wait_until(lambda: c.get(f"/api/runs/{run_id}").json()["status"] == "scraped")
    assert c.post(f"/api/runs/{run_id}/clean", json={}).status_code == 409  # nothing verified yet
    assert c.post(f"/api/runs/{run_id}/verify", json={"brand_set": "acme"}).status_code == 200
    wait_until(lambda: c.get(f"/api/runs/{run_id}").json()["status"] == "verified")
    return run_id


def _done_clean(c, run_id):
    return wait_until(
        lambda: (
            (d := c.get(f"/api/runs/{run_id}").json())["clean_jobs"]
            and d["clean_jobs"][0]["status"] in ("done", "failed")
            and not d["active_job"]
            and d
        )
    )


def test_details_round_trip_and_validation(settings):
    with make_client(settings, search_one=FakeSearch()) as c:
        empty = {"own_websites": [], "own_handles": [], "competitor_websites": [], "competitor_handles": []}
        assert c.get("/api/cleaning-details/acme").json() == {"brand_set": "acme", "saved": False, "details": empty}
        sets = {s["name"]: s for s in c.get("/api/brands").json()["sets"]}
        assert sets["acme"]["cleaning"] == {"details": empty, "saved": False}
        bad = c.put("/api/cleaning-details/acme", json={"own_websites": ["not a site"]})
        assert bad.status_code == 422 and "Own websites" in bad.json()["error"]
        assert c.put("/api/cleaning-details/nope", json={}).status_code == 404
        assert c.get("/api/cleaning-details/nope").status_code == 404
        assert c.put("/api/cleaning-details/acme", json={"extra": []}).status_code == 422
        ok = c.put(
            "/api/cleaning-details/acme",
            json={"own_websites": ["https://www.Acme-Bags.com/about"], "own_handles": ["@acmebags"]},
        )
        assert ok.status_code == 200
        assert ok.json()["details"]["own_websites"] == ["acme-bags.com"]
        # Stored with the brand set in config.yaml, which a hand-written set keeps unchanged otherwise.
        from urlverify.config import load_config

        cfg = load_config(settings.verifier_config)
        assert cfg.cleaning["acme"]["own_handles"] == ["acmebags"]
        assert "acme" in cfg.brands
        got = c.get("/api/cleaning-details/acme").json()
        assert got["saved"] is True and got["details"]["own_handles"] == ["acmebags"]
        sets = {s["name"]: s for s in c.get("/api/brands").json()["sets"]}
        assert sets["acme"]["cleaning"]["details"]["own_websites"] == ["acme-bags.com"]
        # Deleting the brand set deletes its details with it.
        assert c.delete("/api/brands/acme").status_code == 200
        assert "acme" not in load_config(settings.verifier_config).cleaning


def test_form_set_saves_cleaning_with_its_answers(settings):
    profile = {"brands": [{"name": "Zentro", "handles": ["zentro_in"]}]}
    with make_client(settings, search_one=FakeSearch()) as c:
        # A form set without saved details starts from its form's social handles.
        assert c.put("/api/brand-profiles/zentro", json={"profile": profile, "create": True}).status_code == 200
        got = c.get("/api/cleaning-details/zentro").json()
        assert got["saved"] is False and got["details"]["own_handles"] == ["zentro_in"]
        bad = c.put(
            "/api/brand-profiles/zentro", json={"profile": profile, "cleaning": {"competitor_websites": ["x y"]}}
        )
        assert bad.status_code == 422 and "Competitor websites" in bad.json()["error"]
        r = c.put(
            "/api/brand-profiles/zentro",
            json={"profile": profile, "cleaning": {"own_websites": ["zentro.example"], "own_handles": ["zentro_in"]}},
        )
        assert r.status_code == 200
        got = c.get("/api/cleaning-details/zentro").json()
        assert got["saved"] is True and got["details"]["own_websites"] == ["zentro.example"]
        # Saving the answers alone leaves the details as they were.
        assert c.put("/api/brand-profiles/zentro", json={"profile": profile}).status_code == 200
        assert c.get("/api/cleaning-details/zentro").json()["details"]["own_websites"] == ["zentro.example"]


def test_clean_a_verified_run(settings, keys):
    with make_client(settings, search_one=AcmeSearch({"alpha": LINKS}), pipeline_run=fake_pipeline) as c:
        c.put("/api/cleaning-details/acme", json={"own_websites": ["acme-bags.com"], "own_handles": ["acmebags"]})
        run_id = _verified_run(c)
        r = c.post(f"/api/runs/{run_id}/clean")
        assert r.status_code == 200
        # Editing the details after starting does not change this cleaning: it copied them when it started.
        c.put("/api/cleaning-details/acme", json={})
        detail = _done_clean(c, run_id)
        job = detail["clean_jobs"][0]
        assert job["status"] == "done" and job["has_output"]
        assert job["details"]["own_websites"] == ["acme-bags.com"]
        assert detail["status"] == "verified"
        s = job["summary"]
        assert s["rows_in"] == 5
        assert s["brand_communication"] == 2
        assert s["buckets"]["Major Media"] == 1
        assert s["buckets"]["Instagram"] == 1
        assert s["sheets"]["Low Quality"] == 1
        assert s["checks_ok"] is True

        xlsx = c.get(f"/api/runs/{run_id}/clean/{job['id']}/cleaned.xlsx")
        assert xlsx.status_code == 200
        assert "acme_RepScore_clean_2026-03-01_2026-08-31_" in xlsx.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(xlsx.content))
        assert wb.sheetnames[:3] == ["Cleaning Summary", "Clean Data", "Major Media"]
        bc = [str(row[1].value) for row in wb["Brand Communication"].iter_rows(min_row=2)]
        assert sorted(bc) == sorted([LINKS[1], LINKS[3]])

        assert c.get(f"/api/runs/{run_id}/clean/999/cleaned.xlsx").status_code == 404


def test_failed_clean_can_be_retried(settings, keys, monkeypatch):
    import pipeline_api.cleaning as cleaning

    real = cleaning.clean
    calls = {"n": 0}

    def flaky(rows, ctx):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("disk on fire")
        return real(rows, ctx)

    monkeypatch.setattr(cleaning, "clean", flaky)
    with make_client(settings, search_one=AcmeSearch({"alpha": LINKS}), pipeline_run=fake_pipeline) as c:
        run_id = _verified_run(c)
        job_id = c.post(f"/api/runs/{run_id}/clean").json()["job_id"]
        detail = _done_clean(c, run_id)
        assert detail["clean_jobs"][0]["status"] == "failed"
        assert detail["status"] == "failed"
        assert "Cleaning failed: RuntimeError: disk on fire" in detail["error"]
        assert c.get(f"/api/runs/{run_id}/clean/{detail['clean_jobs'][0]['id']}/cleaned.xlsx").status_code == 409
        assert c.post(f"/api/jobs/{job_id}/retry").status_code == 200
        detail = _done_clean(c, run_id)
        assert detail["clean_jobs"][0]["status"] == "done"
        assert detail["status"] == "verified"


def test_clean_rejects_unknown_or_unfinished_verification(settings, keys):
    with make_client(settings, search_one=AcmeSearch({"alpha": LINKS}), pipeline_run=fake_pipeline) as c:
        run_id = _verified_run(c)
        assert c.post(f"/api/runs/{run_id}/clean", json={"verify_job_id": 999}).status_code == 404
        assert c.post("/api/runs/nope/clean").status_code == 404
