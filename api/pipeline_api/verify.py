"""Verify job: verify a run's SERP rows with an immutable brand-rule snapshot, off the API event loop."""

from __future__ import annotations

import asyncio
import contextlib
import json
import math
import sqlite3
import threading
import time
from collections import Counter
from collections.abc import Awaitable, Callable
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from . import runs
from .db import now, session, transaction
from .deps import Deps
from .errors import ApiError
from .export import verified_filename, write_serp_xlsx
from .jobs import JobContext, JobKind, JobOutcome, JobRecord


def snapshot_rules(config_path: Path, brand_set: str) -> str:
    """Resolve the set from a FRESH read of config.yaml and deep-copy it to JSON. Exact name; no fallback."""
    from urlverify.config import load_config

    from .config_errors import CONFIG_LOAD_ERRORS

    try:
        cfg = load_config(config_path)
    except CONFIG_LOAD_ERRORS as e:
        raise ApiError(422, f"The verifier config {config_path} could not be read: {e}") from None
    rules = cfg.brands.get(brand_set)
    if rules is None:
        raise ApiError(
            422,
            f"Brand set {brand_set!r} is not in {config_path.name}. Pick one of: "
            f"{', '.join(sorted(cfg.brands))}, or add it in the brand editor.",
        )
    return json.dumps([asdict(r) for r in rules])


def rules_from_snapshot(snapshot_json: str) -> list[Any]:
    from urlverify.brands import rule_from_dict

    return [rule_from_dict(d) for d in json.loads(snapshot_json)]


class VerifierThread:
    """Runs one pipeline coroutine on its own event loop in a worker thread. urlverify's pipeline does synchronous
    extraction work inside its coroutines; on this loop that can never stall the API's loop."""

    def __init__(self, make_coro: Callable[[], Awaitable[dict[str, Any]]]):
        self._make_coro = make_coro
        self._loop: asyncio.AbstractEventLoop | None = None
        self._task: asyncio.Task[dict[str, Any]] | None = None
        self._started = threading.Event()
        self._cancel_requested = threading.Event()

    def run(self) -> dict[str, Any] | None:
        loop = asyncio.new_event_loop()
        self._loop = loop
        try:
            asyncio.set_event_loop(loop)
            # ensure_future (not create_task) accepts any Awaitable: deps.pipeline_run is typed as one (its real
            # callable is always an async function, so this still wraps an actual coroutine at runtime).
            self._task = asyncio.ensure_future(self._make_coro(), loop=loop)
            self._started.set()
            if self._cancel_requested.is_set():
                self._task.cancel()
            try:
                return loop.run_until_complete(self._task)
            except asyncio.CancelledError:
                if self._cancel_requested.is_set():
                    return None  # a cancel() call asked for exactly this
                # Nothing in this class ever cancels the task on its own, so a CancelledError that arrives without
                # a matching cancel() request is not a cancellation this job recognizes - fail loudly instead of
                # quietly reporting JobOutcome.CANCELLED for something else going wrong.
                raise RuntimeError("verifier was cancelled unexpectedly") from None
        finally:
            self._started.set()
            self._shutdown(loop)

    @staticmethod
    def _shutdown(loop: asyncio.AbstractEventLoop) -> None:
        """Mirrors asyncio.run's own teardown (cancel every remaining task, drain it, then shut down async
        generators and the default executor) before closing the loop - so a task left behind by an early return
        (e.g. the CancelledError branch above) does not leak a pending task/coroutine warning."""
        try:
            pending = asyncio.all_tasks(loop)
            if pending:
                for task in pending:
                    task.cancel()
                loop.run_until_complete(asyncio.gather(*pending, return_exceptions=True))
            loop.run_until_complete(loop.shutdown_asyncgens())
            loop.run_until_complete(loop.shutdown_default_executor())
        finally:
            asyncio.set_event_loop(None)
            loop.close()

    def cancel(self) -> None:
        """Cooperative: takes effect at the pipeline's next await; a synchronous step in progress finishes first."""
        self._cancel_requested.set()
        if self._started.is_set() and self._loop is not None and self._task is not None:
            with contextlib.suppress(RuntimeError):  # loop already closed: nothing left to cancel
                self._loop.call_soon_threadsafe(self._task.cancel)


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, bool | int | str):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) else value
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)


def _read_verified_rows(path: Path) -> list[tuple[int, str | None, int, str]]:
    """Parse the verified sheet into (seq, status, is_duplicate, row_json) tuples. Pure I/O on the sheet only - no
    database access - so a caller can do this parsing before opening a write transaction, and not hold the
    database's write lock while it works through a possibly large sheet."""
    from urlverify.annotate import DUPLICATE_COLUMN

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(it)]
        if "Status" not in header:
            # The job's whole guarantee is "every row carries exactly the verdict the pipeline wrote" - a sheet
            # with no Status column at all means something upstream is badly wrong, not a row to ingest blank.
            raise ValueError(f"verified sheet {path.name} has no 'Status' column (columns: {header})")
        status_i = header.index("Status")
        dup_i = header.index(DUPLICATE_COLUMN) if DUPLICATE_COLUMN in header else None
        batch = []
        for seq, raw in enumerate(it):
            values = list(raw) + [None] * (len(header) - len(raw))
            row = {h: _jsonable(v) for h, v in zip(header, values, strict=True)}
            is_dup = int(dup_i is not None and values[dup_i] not in (None, ""))
            # ensure_ascii=False: keeps literal non-ASCII characters (curly apostrophes, accents, etc.) in the
            # stored JSON so the results endpoint's LIKE search can match them literally - matching
            # runs.store_query_result's own convention.
            batch.append((seq, values[status_i], is_dup, json.dumps(row, ensure_ascii=False)))
        return batch
    finally:
        wb.close()


def _write_verified_rows(
    conn: sqlite3.Connection, verify_job_id: int, rows: list[tuple[int, str | None, int, str]]
) -> int:
    """Replace verify_job_id's rows with `rows`. The caller holds the transaction."""
    conn.execute("DELETE FROM verify_rows WHERE verify_job_id = ?", (verify_job_id,))
    conn.executemany(
        "INSERT INTO verify_rows (verify_job_id, seq, status, is_duplicate, row_json) VALUES (?, ?, ?, ?, ?)",
        [(verify_job_id, seq, status, is_dup, row_json) for seq, status, is_dup, row_json in rows],
    )
    return len(rows)


def ingest_verified(conn: sqlite3.Connection, verify_job_id: int, path: Path) -> int:
    """Copy every row of the verified sheet into verify_rows. The caller holds the transaction."""
    return _write_verified_rows(conn, verify_job_id, _read_verified_rows(path))


class _Progress:
    def __init__(self, total: int):
        self.total = total
        self.done = 0
        self.counts: Counter[str] = Counter()
        self.dirty = False

    def add(self, status: str) -> None:
        self.done += 1
        self.counts[status] += 1
        self.dirty = True


async def run_verify(ctx: JobContext, deps: Deps) -> JobOutcome:
    from urlverify.config import load_config
    from urlverify.load import load_records, read_input

    vj_id = ctx.job.ref_id
    assert vj_id is not None
    with session(ctx.db_path) as conn:
        vj = conn.execute("SELECT * FROM verify_jobs WHERE id = ?", (vj_id,)).fetchone()
        run = runs.get_run(conn, ctx.job.run_id)
        rows = runs.serp_rows(conn, run["id"])
    rules = rules_from_snapshot(vj["brand_rules_json"])  # the job's only brand source
    cfg = load_config(deps.settings.verifier_config)  # fetch, match and verdict settings only
    job_dir = deps.settings.exports_dir / run["id"]
    input_path = job_dir / f"verify-{vj_id}-input.xlsx"
    output_path = job_dir / verified_filename(run, vj_id)
    write_serp_xlsx(input_path, deps.bs, rows)
    total = len(load_records(read_input(input_path), cfg.tracking_params)[0])
    with session(ctx.db_path) as conn:
        conn.execute(
            "UPDATE verify_jobs SET status = 'running', total_urls = ?, done_urls = 0, "
            "status_counts_json = '{}', error = NULL, started_at = ? WHERE id = ?",
            (total, now(), vj_id),
        )
    progress = _Progress(total)
    api_loop = asyncio.get_running_loop()

    def on_result(status: str) -> None:  # called on the verifier thread
        api_loop.call_soon_threadsafe(progress.add, status)

    def flush() -> None:
        if not progress.dirty:
            return
        progress.dirty = False
        counts = dict(progress.counts)
        with session(ctx.db_path) as conn:
            conn.execute(
                "UPDATE verify_jobs SET done_urls = ?, status_counts_json = ? WHERE id = ?",
                (progress.done, json.dumps(counts), vj_id),
            )
        ctx.publish(
            {
                "type": "verify_progress",
                "verify_job_id": vj_id,
                "done_urls": progress.done,
                "total_urls": total,
                "status_counts": counts,
            }
        )

    vt = VerifierThread(
        lambda: deps.pipeline_run(
            input_path,
            output_path,
            cfg,
            vj["brand_set"],
            cache_dir=str(deps.settings.verifier_cache),
            rules=rules,
            on_result=on_result,
        )
    )
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix=f"verify-{vj_id}")
    try:
        fut = api_loop.run_in_executor(executor, vt.run)
        last_flush = time.monotonic()
        while not fut.done():
            if ctx.cancel.is_set():
                vt.cancel()
            await asyncio.wait({fut}, timeout=0.2)
            if time.monotonic() - last_flush >= 1.0:
                flush()
                last_flush = time.monotonic()
        counts = fut.result()
    except (asyncio.CancelledError, Exception) as e:
        # Runner shutdown (CancelledError) or the pipeline itself raising: either way this job did not finish,
        # so neither the per-run input (regenerated from serp_rows on the next attempt) nor a partial output file
        # should be left behind for a later listing/export to trip over.
        if isinstance(e, asyncio.CancelledError):
            vt.cancel()  # stop the thread; the job stays 'running' and is re-queued on restart
        input_path.unlink(missing_ok=True)
        output_path.unlink(missing_ok=True)
        raise
    finally:
        executor.shutdown(wait=False)
    input_path.unlink(missing_ok=True)
    if counts is None:
        output_path.unlink(missing_ok=True)
        return JobOutcome.CANCELLED
    final = {k: v for k, v in counts.items() if not k.startswith("_")}
    # Parse the (possibly large) verified sheet BEFORE opening the write transaction: the immediate lock is then
    # held only for the few small writes that follow, not for however long parsing takes.
    parsed_rows = _read_verified_rows(output_path)
    with session(ctx.db_path) as conn, transaction(conn, immediate=True):
        _write_verified_rows(conn, vj_id, parsed_rows)
        conn.execute(
            "UPDATE verify_jobs SET status = 'done', done_urls = ?, status_counts_json = ?, output_path = ?, "
            "finished_at = ? WHERE id = ?",
            (counts.get("_unique", total), json.dumps(final), str(output_path), now(), vj_id),
        )
        runs.set_run_status(conn, run["id"], "verified")
    ctx.publish(
        {
            "type": "verify_progress",
            "verify_job_id": vj_id,
            "done_urls": counts.get("_unique", total),
            "total_urls": total,
            "status_counts": final,
        }
    )
    return JobOutcome.DONE


def verify_kind(deps: Deps) -> JobKind:
    async def run(ctx: JobContext) -> JobOutcome:
        return await run_verify(ctx, deps)

    def on_failed(conn: sqlite3.Connection, job: JobRecord, message: str) -> None:
        conn.execute(
            "UPDATE verify_jobs SET status = 'failed', error = ?, finished_at = ? WHERE id = ?",
            (message, now(), job.ref_id),
        )
        runs.set_run_status(
            conn,
            job.run_id,
            "failed",
            f"Verification failed: {message}. Use Retry to run it again with the same brand snapshot.",
        )

    def on_cancelled(conn: sqlite3.Connection, job: JobRecord) -> None:
        conn.execute("UPDATE verify_jobs SET status = 'cancelled', finished_at = ? WHERE id = ?", (now(), job.ref_id))
        runs.set_run_status(conn, job.run_id, runs.after_stop_status(conn, job.run_id))

    return JobKind(run=run, on_failed=on_failed, on_cancelled=on_cancelled)
