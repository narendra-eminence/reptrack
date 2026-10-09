"""Map links to Major / Regional / Other Media with the RepScore Master Media List.

Ported from the repscore-data-processing skill's scripts/media_bucket.py; keep the two in step.

Matching order (first hit wins):
  1. exact host + longest Website path prefix         (business-standard.com/markets -> BS Markets)
  2. exact host, bare Website                          (jagran.com -> Dainik Jagran, the parent row)
  3. parent domain of the host, never the public suffix (epaper.jagran.com -> jagran.com)
  4. publication name or alias == the export's source name (rows with no usable Website)
  5. no match -> Other Media, status NEW, Level N/A, Type blank

When several master rows share a bare host and none has a path, the parent masthead wins: rows whose Parent / Group
is another of the sharing rows (ET Foods, parent Economic Times) are set aside, then the shortest Publication name
wins, unless the export's source name equals one of the sharing rows' Publication or Alias.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl

SHEET = "Master Media List"
BUCKET_LABEL = {"Major": "Major Media", "Regional": "Regional Media", "Other": "Other Media"}
REQUIRED_COLUMNS = (
    "Media ID",
    "Bucket",
    "Level",
    "Publication",
    "Type",
    "Geography",
    "Region / State",
    "Language",
    "Format",
    "Website",
    "Parent / Group",
    "Alias / Variant Names",
)

# Two-part public suffixes seen in Indian media domains. A host is never reduced past these.
TWO_PART_SUFFIXES = {
    "co.in",
    "com.au",
    "co.uk",
    "org.in",
    "net.in",
    "gov.in",
    "ac.in",
    "edu.in",
    "nic.in",
    "com.sg",
    "co.nz",
    "com.pk",
    "com.bd",
    "com.np",
    "co.jp",
    "com.hk",
}
STRIP_LEADING = ("www.", "m.", "amp.", "mobile.", "wap.")
_SCHEME = re.compile(r"^[a-z][a-z0-9+.-]*://", re.I)


class MasterListError(ValueError):
    pass


def _with_scheme(url: str) -> str:
    u = str(url or "").strip()
    return u if _SCHEME.match(u) else "http://" + u


def clean_host(url: str) -> str:
    """The host of a URL, lowercased, without www./m./amp. prefixes or a port."""
    if not url:
        return ""
    try:
        host = (urlsplit(_with_scheme(url)).hostname or "").lower().rstrip(".")
    except ValueError:  # e.g. an unbalanced IPv6 bracket in a scraped link
        return ""
    changed = True
    while changed:
        changed = False
        for p in STRIP_LEADING:
            if host.startswith(p) and host.count(".") >= 2:
                host = host[len(p) :]
                changed = True
    return host


def clean_path(url: str) -> str:
    try:
        return urlsplit(_with_scheme(url)).path.lower().rstrip("/")
    except ValueError:
        return ""


def registrable_domain(host: str) -> str:
    """The shortest registrable form of host: domain + public suffix."""
    parts = host.split(".")
    if len(parts) <= 2:
        return host
    if ".".join(parts[-2:]) in TWO_PART_SUFFIXES:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def parent_hosts(host: str) -> Iterator[str]:
    """host, then each parent up to and including the registrable domain."""
    stop = registrable_domain(host)
    cur = host
    while True:
        yield cur
        if cur == stop or "." not in cur:
            return
        cur = cur.split(".", 1)[1]


def norm_name(s: object) -> str:
    text = str(s or "").lower().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


@dataclass
class Row:
    media_id: str
    bucket: str  # Major | Regional | Other  (raw master value)
    level: str  # Core | Level 1 | Level 2 | N/A
    publication: str
    type: str
    geography: str
    region: str
    language: str
    format: str
    website: str  # raw cell value, may be blank or a non-URL
    parent: str
    alias: str
    host: str = ""  # cleaned host from website, "" when website is unusable
    path: str = ""  # cleaned path from website

    @property
    def bucket_label(self) -> str:
        return BUCKET_LABEL.get(self.bucket, "Other Media")


@dataclass
class Hit:
    bucket: str
    status: str  # Known | NEW
    matched_by: str  # domain+path | domain | parent-domain | name | none
    row: Row | None = None
    domain: str = ""
    notes: list[str] = field(default_factory=list)
    review: bool = False  # the export's source name disagrees with the match: a person should look

    @property
    def level(self) -> str:
        return (self.row.level or "N/A") if self.row else "N/A"

    @property
    def type(self) -> str:
        return self.row.type if self.row else ""

    @property
    def publication(self) -> str:
        return self.row.publication if self.row else ""

    @property
    def media_id(self) -> str:
        return self.row.media_id if self.row else ""

    @property
    def known_by_domain(self) -> bool:
        return self.status == "Known" and self.matched_by != "name"


class MasterMediaList:
    def __init__(self, xlsx: Path | str):
        self.path = Path(xlsx)
        self.rows: list[Row] = []
        self.by_host: dict[str, list[Row]] = {}
        self.by_name: dict[str, list[Row]] = {}
        self.problems: list[str] = []
        self._load()

    def _load(self) -> None:
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        try:
            if SHEET not in wb.sheetnames:
                raise MasterListError(f"sheet {SHEET!r} not found in {self.path.name}")
            it = wb[SHEET].iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(it, ())]
            ix = {h: n for n, h in enumerate(header)}
            missing = [h for h in REQUIRED_COLUMNS if h not in ix]
            if missing:
                raise MasterListError(f"{self.path.name} is missing columns {missing}")

            for r in it:
                if not any(v not in (None, "") for v in r):
                    continue

                def cell(h: str, r: tuple = r) -> str:
                    v = r[ix[h]] if ix[h] < len(r) else None
                    return "" if v is None else str(v).strip()

                row = Row(*(cell(h) for h in REQUIRED_COLUMNS))
                self._check(row)
                self.rows.append(row)
                if row.host:
                    self.by_host.setdefault(row.host, []).append(row)
                for nm in (row.publication, row.alias):
                    if nm:
                        self.by_name.setdefault(norm_name(nm), []).append(row)
        finally:
            wb.close()
        for host, rows in self.by_host.items():
            if host == "google.com" and any(not r.path for r in rows):
                ids = ", ".join(r.media_id for r in rows if not r.path)
                self.problems.append(f"bare shared-platform Website '{host}' on {ids}; it will claim every subdomain")

    def _check(self, row: Row) -> None:
        if row.bucket not in BUCKET_LABEL:
            self.problems.append(f"{row.media_id}: Bucket '{row.bucket}' is not Major/Regional/Other")
        if row.bucket == "Major" and row.level not in ("Core", "Level 1", "Level 2"):
            self.problems.append(f"{row.media_id}: Major row without a Level ('{row.level}')")
        if row.bucket != "Major" and row.level not in ("N/A", ""):
            self.problems.append(f"{row.media_id}: {row.bucket} row carries Level '{row.level}'")
        if row.website:
            if re.match(r"^(https?://)?[a-z0-9-]+(\.[a-z0-9-]+)+", row.website, re.I):
                row.host = clean_host(row.website)
                row.path = clean_path(row.website)
            else:
                self.problems.append(
                    f"{row.media_id} ({row.publication}): Website cell holds '{row.website}', not a URL; "
                    "row is name-match only"
                )

    def _pick(self, rows: list[Row], path: str, source_name: str) -> tuple[Row, str]:
        """Choose one row among rows sharing a host."""
        with_path = [r for r in rows if r.path and (path == r.path or path.startswith(r.path + "/"))]
        if with_path:
            return max(with_path, key=lambda r: len(r.path)), "domain+path"
        bare = [r for r in rows if not r.path] or rows  # every row has a path and none matched: parent masthead
        if len(bare) == 1:
            return bare[0], "domain"
        sn = norm_name(source_name)
        if sn:
            named = [r for r in bare if sn in (norm_name(r.publication), norm_name(r.alias))]
            if named:
                return named[0], "domain"
        pubs = {norm_name(r.publication) for r in bare}
        mastheads = [r for r in bare if norm_name(r.parent) not in pubs - {norm_name(r.publication)}] or bare
        return min(mastheads, key=lambda r: (len(r.publication), r.media_id)), "domain"

    def match(self, link: str, source_name: str = "") -> Hit:
        host = clean_host(link)
        path = clean_path(link) if host else ""
        sn = norm_name(source_name)
        if host:
            for depth, h in enumerate(parent_hosts(host)):
                rows = self.by_host.get(h)
                if not rows:
                    continue
                row, how = self._pick(rows, path, source_name)
                if depth and how == "domain":
                    how = "parent-domain"
                hit = Hit(row.bucket_label, "Known", how, row, registrable_domain(host))
                if len(rows) > 1 and how != "domain+path":
                    hit.notes.append(f"{h} is shared by {len(rows)} master rows; picked {row.publication}")
                if sn and sn not in (norm_name(row.publication), norm_name(row.alias)):
                    other = [x for x in self.by_name.get(sn, []) if x is not row]
                    if other:
                        o = other[0]
                        hit.notes.append(
                            f"source name '{source_name}' is master row {o.media_id} ({o.bucket_label}); "
                            "domain match kept, review"
                        )
                        hit.review = True
                return hit
        if sn and sn in self.by_name:
            rows = self.by_name[sn]
            row = rows[0]
            hit = Hit(row.bucket_label, "Known", "name", row, registrable_domain(host) if host else "")
            if len(rows) > 1:
                hit.notes.append(f"name '{source_name}' matches {len(rows)} master rows; picked {row.media_id}")
                hit.review = True
            if host and row.host and registrable_domain(row.host) != registrable_domain(host):
                hit.notes.append(f"name matched {row.publication} but link host {host} differs from master {row.host}")
                hit.review = True
            return hit
        return Hit("Other Media", "NEW", "none", None, registrable_domain(host) if host else "")

    def counts(self) -> dict[str, int]:
        out = {"rows": len(self.rows)}
        for r in self.rows:
            label = BUCKET_LABEL.get(r.bucket, "Other Media")
            out[label] = out.get(label, 0) + 1
        return out


_cache: dict[tuple[str, float, int], MasterMediaList] = {}


def load_master_list(path: Path) -> MasterMediaList:
    """The parsed list, re-read only when the file on disk changes (replacing the xlsx needs no restart)."""
    st = path.stat()
    key = (str(path.resolve()), st.st_mtime, st.st_size)
    if key not in _cache:
        _cache.clear()
        _cache[key] = MasterMediaList(path)
    return _cache[key]
