"""Write a CleanResult as the formatted RepScore cleaning workbook."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..export import save_atomic, set_cell
from .engine import CleanResult, Sheet

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="E7E6E6")
NEW_FILL = PatternFill("solid", fgColor="FFEB9C")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
WIDTHS = {
    "Title": 55,
    "Opening Text": 60,
    "Hit Sentence": 60,
    "Verified Context": 60,
    "Sentiment Rationale": 55,
    "Insight": 55,
    "Script Hit Sentence": 55,
    "Notes": 45,
    "Link": 45,
    "Example link": 45,
    "Query": 40,
    "Destination": 40,
    "Cleaning Outcome": 34,
    "Snippet": 60,
}
WRAPPED = {"Title", "Opening Text", "Hit Sentence", "Verified Context", "Sentiment Rationale", "Insight", "Notes"}


def _header(ws, columns: list[str]) -> None:
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS.get(name, min(max(len(name) + 4, 12), 26))
    ws.freeze_panes = "A2"


def _table(ws, sheet: Sheet) -> None:
    _header(ws, sheet.columns)
    wrap_cols = [n for n, c in enumerate(sheet.columns, start=1) if c in WRAPPED]
    status_col = sheet.columns.index("Status") if "Status" in sheet.columns else None
    for r, row in enumerate(sheet.rows, start=2):
        is_new = status_col is not None and row.get("Status") == "NEW"
        for col, name in enumerate(sheet.columns, start=1):
            cell = set_cell(ws, r, col, row.get(name))
            if isinstance(cell.value, dt.date):
                cell.number_format = "yyyy-mm-dd"
            if is_new:
                cell.fill = NEW_FILL
        for col in wrap_cols:
            ws.cell(row=r, column=col).alignment = WRAP
    last = get_column_letter(len(sheet.columns))
    ws.auto_filter.ref = f"A1:{last}{max(1, len(sheet.rows) + 1)}"


def _summary(ws, sheet: Sheet) -> None:
    _header(ws, sheet.columns)
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 90
    for r, row in enumerate(sheet.rows, start=2):
        a = set_cell(ws, r, 1, row["Item"])
        b = set_cell(ws, r, 2, row["Value"])
        b.alignment = WRAP if isinstance(row["Value"], str) else TOP
        if row.get("_section"):
            a.font = Font(bold=True)
            a.fill = b.fill = SECTION_FILL


def write_workbook(path: Path, result: CleanResult) -> None:
    wb = Workbook()
    first = wb.active
    assert first is not None
    wb.remove(first)
    for sheet in result.sheets:
        ws = wb.create_sheet(sheet.name[:31])
        if sheet.kind == "summary":
            _summary(ws, sheet)
        else:
            _table(ws, sheet)
    save_atomic(wb, path)
